#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把 Markdown 源文件转成 Word 文档（支持 # 标题、- 列表、1. 列表、| 表格 |、**粗体**段落）。

用法： python3 build_docx.py 源.md 输出.docx "文档标题"
"""
import sys, re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


def set_cn_font(doc, name="宋体", size=10.5):
    st = doc.styles["Normal"]
    st.font.name = name
    st.font.size = Pt(size)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), name)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.25


def add_runs(p, text):
    """处理 **粗体** 与 `代码`"""
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = p.add_run(part[1:-1]); r.italic = True
        elif part.startswith("`") and part.endswith("`"):
            r = p.add_run(part[1:-1]); r.font.name = "Menlo"; r.font.size = Pt(9)
        else:
            p.add_run(part)


def build(src, out, title=None):
    lines = open(src, encoding="utf-8").read().split("\n")
    doc = Document()
    set_cn_font(doc)
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(2.2)
        s.left_margin = s.right_margin = Cm(2.4)

    if title:
        h = doc.add_heading(title, level=0)
        for r in h.runs:
            r.font.name = "微软雅黑"
            r.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln.strip():
            i += 1; continue

        # 表格
        if ln.lstrip().startswith("|") and i + 1 < len(lines) and set(lines[i+1].replace("|", "").strip()) <= set("-: "):
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not set("".join(cells)) <= set("-: "):
                    rows.append(cells)
                i += 1
            ncol = max(len(r) for r in rows)
            t = doc.add_table(rows=0, cols=ncol)
            t.style = "Light Grid Accent 1"
            for ri, r in enumerate(rows):
                cells = t.add_row().cells
                for ci in range(ncol):
                    txt = r[ci] if ci < len(r) else ""
                    para = cells[ci].paragraphs[0]
                    add_runs(para, txt)
                    for run in para.runs:
                        run.font.size = Pt(9)
                        if ri == 0:
                            run.bold = True
            doc.add_paragraph()
            continue

        m = re.match(r"^>\s?(.*)$", ln)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            add_runs(p, m.group(1))
            for r in p.runs:
                r.italic = True
            i += 1; continue

        m = re.match(r"^(#{1,6})\s+(.*)$", ln)
        if m:
            lvl = len(m.group(1))
            h = doc.add_heading(m.group(2), level=min(lvl, 4))
            for r in h.runs:
                r.font.name = "微软雅黑"
                r.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
                if lvl <= 2:
                    r.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
            i += 1; continue

        m = re.match(r"^\s*[-*+]\s+(.*)$", ln)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            add_runs(p, m.group(1)); i += 1; continue

        m = re.match(r"^\s*(\d+)[.、)]\s+(.*)$", ln)
        if m:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(2)
            add_runs(p, m.group(2)); i += 1; continue

        if ln.strip() in ("---", "***", "___"):
            i += 1; continue

        p = doc.add_paragraph()
        add_runs(p, ln)
        i += 1

    doc.save(out)
    print("已生成:", out)


if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2], sys.argv[3] if len(sys.argv) > 3 else None)
