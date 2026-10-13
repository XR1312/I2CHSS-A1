#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""在《4_AI论文阅读工具评估报告.docx》末尾追加"附录二"（2026-10-11 补记）。"""
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn

DOC = '/Users/XR1312/Desktop/SC-1/完成的作业/4_AI论文阅读工具评估报告.docx'
HEAD = '附录二：四条只做文本核对查不出来的问题（2026-10-11 补记）'

TITLE = '附录二：四条只做文本核对查不出来的问题（2026-10-11 补记）'

LEAD = ('第 2 部分的 5 次运行里，工具在"核对我点名的数字"这件事上表现很好：表 C2 的 50 余个数值逐格正确。'
        '但当我按论文 Data availability 声明的指引，把 OSF 项目（osf.io/atj5u）里公开的 69×15 输入矩阵与官方 R 脚本取来、'
        '把潜类别分析重跑了一遍之后，又发现四条它完全没有触及的问题。这四条都不属于"补充材料里写着、工具该抄没抄"，'
        '而属于"把数字拿进公式或模型算一遍才会暴露"的那一层，因此正好用来标出文本核对的能力边界。'
        '需要说明：运行时的输入只有正文与补充材料，工具手上没有脚本，所以第 4 条无从查起；但第 1、2 条**仅凭论文自身就能发现**，'
        '第 3 条也只需多做一步稳定性检查，它们是真实的漏检。')

ITEMS = [
    ('1. 表 C2 中名为"AIC"的那一列，不可能是 AIC。',
     '这一列每个 K 都比同行的 BIC 大（K=1：1380 > 1365；K=2：979 > 948；K=3：998 > 951）。'
     '但 AIC = −2ll + 2k，BIC = −2ll + k·ln N，本例 N = 69、ln N ≈ 4.23 > 2，所以 AIC 必然恒小于 BIC，'
     '不可能反过来。工具逐格核对了这一列的数字（并且核对无误），却没有做这个不等式检查。'
     '取来脚本即可看到原因：脚本写的是 `data.frame(..., AIC = CAICs, ...)`，而 `CAICs <- (-2*lc$llik) + lc$npar * (1 + log(lc$N))`'
     '——这是 CAIC（一致赤池信息准则）。真正的 AIC 在 K=5 才取到最小，K=3 时是 846.0。'
     '这是一个"列名与公式不符"的问题，也是一处只要多问一句"这个数算得出来吗"就能抓住的漏检。'),
    ('2. 发现"BIC 规则不成立"之后，没有再往下追一步。',
     '工具最亮眼的一次发现是指出：按脚本自己的规则（取 BIC 最小）应当选 2 类解，作者却报告了 3 类解。'
     '但它停在"作者没照自己的规则做"，没有继续核对整套准则的方向——BIC 与 CAIC 指向 2 类，aBIC 一路降到 10 类，'
     '真 AIC 指向 5 类，三个常用准则没有一个支持 3 类。'
     '"类别数是数据决定的还是解释力决定的"是这篇论文最可质疑的地方，而结论其实就摆在它已经读过的那张表里。'),
    ('3. 把局部最优解当成了可信基准。',
     '我把自己实现的 EM 随机起点数从 30 加到 400，K=6、K=8、K=9 的对数似然分别还能再改善 3.7、3.8、4.3。'
     '也就是说表 C2 里 K≥5 各行是起点依赖的局部极大值，那一带的 BIC/aBIC 不能横向比较（aBIC 的单调下降也因此不可信）。'
     '工具把这张表当成可信基准逐格核对，没有质疑"这些数是不是该模型的最优解"。'
     '这一步不需要脚本也能想得到：K=3 时自由参数已有 47 个、样本只有 69 个，似然面多峰是常识性风险。'),
    ('4. 15 条标准里有 3 条取值完全相同。',
     '公开矩阵中 `incl_gender`、`incl_age`、`incl_educ` 三列完全一致（各 34 个 1，两两相关系数 1.000），'
     '也就是"性别包容性/年龄包容性/教育包容性"三个指标在 69 个域名上给出了同一串 0/1。'
     '因此进入模型的信息实际只有 13 条互异的标准，而脚本仍按 15 条计算自由参数（npar = K·15 + (K−1)），'
     'K=3 时的 47 个自由参数是偏高的。工具（以及我最初的阅读）都把"六维度 15 条标准"当成 15 个独立指标接受了。'
     '这一条要读到数据文件才可能发现，可以算作输入缺失导致的漏检，但它提示了一个更普遍的问题：'
     '工具的"核查"止于被喂进去的文本，不会主动去取来数据看看这些列到底长什么样。'),
]

TAIL = ('把四条放在一起看，共同点是：工具的审计深度取决于我有没有要求它核算，而且即便要求了，它核算的也只是我递进去的数字，'
        '不会主动去把数据与脚本取来、重建一遍计算过程。这一步对复现类任务恰好是最关键的一步，目前仍然要人来做。'
        '这也从反面支持了本报告第 3 节提炼的"提问配方"：要它查什么，就得先把材料递到它手上，并把"算一遍、比一遍"明确写进要求里。'
        '四条的证据与可重跑的脚本见 `附录_代码与数据/`（`repro_lca.py`、`repro_report.txt`、'
        '`OSF数据可得性核查记录.txt`）；本报告与《论文阅读工具输出记录》的其余部分均未改动。')


def add_runs(p, text):
    for part in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith('`') and part.endswith('`'):
            r = p.add_run(part[1:-1]); r.font.name = 'Menlo'; r.font.size = Pt(9)
        else:
            p.add_run(part)


doc = Document(DOC)
assert not any(p.text.strip() == TITLE for p in doc.paragraphs), 'already appended'

h = doc.add_heading(TITLE, level=2)
for r in h.runs:
    r.font.name = '微软雅黑'
    r.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    r.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)

p = doc.add_paragraph(); add_runs(p, LEAD)

for head, body in ITEMS:
    hp = doc.add_paragraph(style='List Number')
    hp.paragraph_format.space_after = Pt(2)
    add_runs(hp, '**' + head + '**')
    bp = doc.add_paragraph()
    bp.paragraph_format.left_indent = Cm(0.6)
    add_runs(bp, body)

p = doc.add_paragraph(); add_runs(p, TAIL)
doc.save(DOC)
print('appended ->', DOC)
